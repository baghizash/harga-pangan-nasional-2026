#!/usr/bin/env python3
"""Build Portofolio-Analisis-Harga-Pangan.pptx (16:9, 10 slides, ID).
REDESIGN design-taste: tema T3 (oat, plum, tangerine), varian Soft Premium.
Konten, urutan slide, dan semua angka identik dengan deck lama
(sumber: build_ppt_legacy.py + findings.json, data PIHPS/BI).
Sistem visual: sudut rounded konsisten, bento asimetris, double-bezel,
plum untuk struktur, tangerine untuk penekanan. Tanpa gradien/shadow.
Dial: VARIANCE 7, MOTION 5, DENSITY 3. Nol em-dash/en-dash.
"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
F = json.loads((HERE / "findings.json").read_text())
CHARTS = HERE / "charts"
OUT = Path.home() / "workspace" / "your_files" / "Portofolio-Analisis-Harga-Pangan.pptx"

# ---- Palet T3 terkunci ----
GROUND = RGBColor(0xF0, 0xE8, 0xD8)    # oat
SURFACE = RGBColor(0xF8, 0xF2, 0xE4)   # permukaan oat terang
INK = RGBColor(0x2A, 0x1E, 0x33)       # tinta plum gelap
PLUM = RGBColor(0x5E, 0x2A, 0x72)      # primary
TANG = RGBColor(0xEF, 0x7D, 0x1A)      # aksen tangerine
PLUM_L = RGBColor(0x8A, 0x5C, 0xA6)    # plum muda
BEZEL = RGBColor(0xDD, 0xCF, 0xB4)    # bingkai luar bezel (oat lebih dalam)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Noto Sans"
REPO_URL = "github.com/baghizash/harga-pangan-nasional-2026"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
SW, SH, M = 13.333, 7.5, 0.55


def rp(v):
    return f"Rp{v:,}".replace(",", ".")


def slide_new():
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = GROUND
    return s


def _noline(sh):
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def rect(s, x, y, w, h, color):
    sh = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    return _noline(sh)


def rrect(s, x, y, w, h, color, radius=0.16):
    sh = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sh.adjustments[0] = max(0.0, min(0.5, radius / min(w, h)))
    except Exception:
        pass
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    return _noline(sh)


def bezel(s, x, y, w, h, inner=SURFACE, outer=BEZEL, pad=0.07, radius=0.18):
    rrect(s, x, y, w, h, outer, radius + pad)
    return rrect(s, x + pad, y + pad, w - 2 * pad, h - 2 * pad, inner, radius)


def tb(s, x, y, w, h):
    tx = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def par(tf, text, size, color, bold=False, italic=False, first=False,
        align=None, after=6, before=0, leading=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(after)
    p.space_before = Pt(before)
    p.line_spacing = leading
    if align is not None:
        p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = FONT
    return p


def bullets(tf, items, size=15.5, gap=12, leading=1.12, color=INK):
    for i, t in enumerate(items):
        p = par(tf, "", size, color, first=(i == 0), after=gap, leading=leading)
        pPr = p._p.get_or_add_pPr()
        bu = pPr.makeelement(qn("a:buChar"), {"char": "\u2022"})
        pPr.append(bu)
        pPr.set("marL", "190500"); pPr.set("indent", "-190500")
        r = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = FONT
    return tf


def title(s, text, y=0.42):
    tf = tb(s, M, y, 12.2, 1.0)
    par(tf, text, 31, INK, bold=True, first=True, after=0)
    rrect(s, M + 0.02, y + 0.78, 0.85, 0.09, TANG, radius=0.045)


def stat_block(s, x, y, w, big, label, big_color=TANG, big_size=54,
               label_color=None, align=PP_ALIGN.LEFT):
    tf = tb(s, x, y, w, 1.9)
    par(tf, big, big_size, big_color, bold=True, first=True, after=2)
    par(tf, label, 13, label_color or INK, after=0, leading=1.1)


kl = F["kelompok"]; km = F["komoditas"]; pv = F["provinsi"]; lk = F["lintas_komoditas"]
MG = kl["Minyak Goreng"]

# ================= Slide 1: Cover (Asymmetric Split + Data Hero) =================
s = slide_new()
tf = tb(s, M, 0.95, 6.9, 2.2)
par(tf, "Analisis Harga Pangan", 38, INK, bold=True, first=True, after=0, leading=1.05)
par(tf, "Nasional Indonesia 2026", 38, PLUM, bold=True, after=0, leading=1.05)
tf = tb(s, M, 2.85, 6.4, 1.3)
par(tf, "Beras, daging sapi, minyak goreng, gula pasir: tren, volatilitas, "
        "dan disparitas provinsi.", 17.5, INK, first=True, after=0, leading=1.2)
tf = tb(s, M, 4.5, 6.5, 1.0)
par(tf, "Data harian PIHPS / Bank Indonesia: 13 komoditas, 34 provinsi, "
        "1 Januari 2025 sampai 23 Oktober 2026.", 13, INK, first=True, after=0, leading=1.25)
tf = tb(s, M, 6.55, 6.5, 0.5)
par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 14, PLUM, bold=True, first=True, after=0)
# Panel hero kanan: double-bezel plum
bezel(s, 7.45, 0.85, 5.35, 4.55, inner=PLUM, outer=PLUM_L, pad=0.08, radius=0.2)
tf = tb(s, 7.95, 1.35, 4.4, 3.8)
par(tf, "KENAIKAN TERTINGGI 2026", 13, SURFACE, bold=True, first=True, after=6)
par(tf, "+6,8%", 76, TANG, bold=True, after=4)
par(tf, "Minyak Goreng", 21, WHITE, bold=True, after=6)
par(tf, f"Rata-rata {rp(MG['rata2025'])} menjadi {rp(MG['rata2026'])} per kg, "
        "tertinggi dari empat kelompok bahan pokok.", 13.5, SURFACE, after=0, leading=1.25)
# Chip Power BI
rrect(s, 7.45, 5.72, 5.35, 0.78, TANG, radius=0.14)
tf = tb(s, 7.75, 5.86, 4.85, 0.55)
par(tf, "Plus paket dashboard Power BI: star schema, DAX, panduan rakit.",
    13, WHITE, bold=True, first=True, after=0, leading=1.15)

# ================= Slide 2: Kenapa Harga Pangan? =================
s = slide_new()
title(s, "Kenapa Harga Pangan?")
bezel(s, M, 1.55, 7.55, 4.6)
tf = tb(s, 0.95, 1.95, 6.8, 3.9)
bullets(tf, [
    "Beras, daging sapi, minyak goreng, dan gula pasir adalah bahan pokok yang "
    "harganya langsung menentukan daya beli jutaan rumah tangga Indonesia.",
    "Ketika harga pangan naik 5-7% dalam setahun, yang paling terpukul adalah "
    "keluarga berpendapatan rendah, porsi belanja pangannya paling besar.",
    "Analisis ini menjawab: kelompok mana yang naik paling tinggi di 2026, "
    "komoditas mana yang paling tidak stabil, di provinsi mana harganya paling mahal, "
    "dan apakah Ramadan benar-benar mendongkrak harga.",
    "Semua jawaban diambil dari data harian resmi PIHPS/BI, bukan perkiraan.",
], size=15.5, gap=13)
# Rel stat kanan
rrect(s, 8.45, 1.55, 4.33, 1.32, SURFACE, radius=0.16)
stat_block(s, 8.8, 1.73, 3.7, "13", "komoditas kebutuhan pokok dianalisis", big_size=34)
rrect(s, 8.45, 3.03, 4.33, 1.32, SURFACE, radius=0.16)
stat_block(s, 8.8, 3.21, 3.7, "34", "provinsi dibandingkan harganya", big_size=34)
rrect(s, 8.45, 4.51, 4.33, 1.64, PLUM, radius=0.16)
stat_block(s, 8.8, 4.71, 3.7, "5-7%", "kenaikan harga pangan dalam setahun yang "
           "paling terasa bagi keluarga berpendapatan rendah",
           big_size=34, label_color=SURFACE)
tf = tb(s, M, 6.42, 12.2, 0.45)
par(tf, "Sumber data: API publik PIHPS / Bank Indonesia (bi.go.id/hargapangan), Pasar Tradisional.",
    11.5, PLUM_L, italic=True, first=True, after=0)

# ================= Slide 3: Data dan Metode =================
s = slide_new()
title(s, "Data dan Metode")
# Rel stat kiri
rrect(s, M, 1.55, 3.6, 2.2, PLUM, radius=0.16)
stat_block(s, 0.9, 1.82, 2.95, "204.809", "baris data provinsi dari API PIHPS",
           big_size=30, label_color=SURFACE)
rrect(s, M, 3.95, 3.6, 2.2, SURFACE, radius=0.16)
stat_block(s, 0.9, 4.22, 2.95, "6.149", "baris data nasional, disimpan terpisah "
           "dari level provinsi", big_size=30)
bezel(s, 4.5, 1.55, 8.28, 4.6)
tf = tb(s, 4.9, 1.92, 7.5, 3.95)
bullets(tf, [
    "Data harian API PIHPS: 1 Jan 2025 - 23 Okt 2026 (473 hari laporan), "
    "6.149 baris nasional + 204.809 baris provinsi.",
    "13 komoditas daun: 6 kualitas beras, 2 daging sapi, 3 minyak goreng, 2 gula pasir. "
    "API ini tidak menyediakan seri daun untuk telur/daging ayam/bawang, tidak dipakai, tidak dikarang.",
    "Level nasional dan provinsi disimpan di tabel fakta TERPISAH sejak awal "
    "(fact_nasional, fact_provinsi), pelajaran dari proyek cabai, sehingga agregasi "
    "dan measure DAX selalu membaca satu grain tanpa guard khusus.",
    "9 kombinasi komoditas-provinsi tidak memiliki seri laporan (mis. beras di Gorontalo), "
    "dicatat sebagai gap, tidak diisi angka lain.",
], size=15, gap=12)

# ================= Slide 4: Tren Nasional =================
s = slide_new()
title(s, "Tren Nasional 2025-2026")
tf = tb(s, M, 1.6, 4.95, 3.6)
bullets(tf, [
    f"Minyak Goreng rata-rata 2026 {rp(MG['rata2026'])}/kg "
    f"({MG['yoy_pct']:+}% vs 2025): kenaikan tertinggi.",
    f"Daging Sapi {rp(kl['Daging Sapi']['rata2026'])}/kg "
    f"({kl['Daging Sapi']['yoy_pct']:+}%); Beras {rp(kl['Beras']['rata2026'])}/kg "
    f"({kl['Beras']['yoy_pct']:+}%); Gula Pasir {rp(kl['Gula Pasir']['rata2026'])}/kg "
    f"({kl['Gula Pasir']['yoy_pct']:+}%).",
    "Keempat kelompok bergerak naik sepanjang 2026, tekanan harga pangan bersifat luas, "
    "bukan satu komoditas saja.",
], size=15.5, gap=13)
rrect(s, M, 5.3, 4.95, 1.28, PLUM, radius=0.15)
stat_block(s, 0.85, 5.5, 4.4, rp(MG["rata2026"]), "rata-rata minyak goreng per kg, 2026",
           big_size=27, label_color=SURFACE)
bezel(s, 5.85, 1.55, 6.93, 4.35)
s.shapes.add_picture(str(CHARTS / "01_tren_nasional.png"), Inches(6.12), Inches(1.98),
                     width=Inches(6.4))

# ================= Slide 5: Siapa Naik Paling Tinggi? =================
s = slide_new()
title(s, "Siapa Naik Paling Tinggi?")
bezel(s, M, 1.55, 5.6, 3.85)
s.shapes.add_picture(str(CHARTS / "02_yoy_kelompok.png"), Inches(0.88), Inches(1.95),
                     width=Inches(4.95))
tf = tb(s, 6.55, 1.62, 6.25, 3.8)
bullets(tf, [
    f"Minyak Goreng: +{MG['yoy_pct']}%, juara kenaikan, "
    f"didorong Minyak Goreng Curah (+{km['com_17']['yoy_pct']}%).",
    f"Daging Sapi: +{kl['Daging Sapi']['yoy_pct']}%, kedua tertinggi, "
    f"harga rata-rata 2026 menembus {rp(kl['Daging Sapi']['rata2026'])}/kg.",
    f"Beras: +{kl['Beras']['yoy_pct']}%, moderat tapi basisnya besar: "
    f"Rp15.515 ke {rp(kl['Beras']['rata2026'])}/kg.",
    f"Gula Pasir: +{kl['Gula Pasir']['yoy_pct']}%, kenaikan paling rendah, "
    "harga paling terkendali.",
], size=15.5, gap=12)
rrect(s, M, 5.68, 12.23, 1.02, SURFACE, radius=0.14)
tf = tb(s, 1.05, 5.9, 11.4, 0.6)
par(tf, "Kenaikan 2026 vs 2025:   Minyak Goreng +6,8%      Daging Sapi +6,0%      "
        "Beras +3,1%      Gula Pasir +2,4%", 16, PLUM, bold=True, first=True, after=0)

# ================= Slide 6: Volatilitas =================
s = slide_new()
title(s, "Volatilitas: Mana yang Paling Liar?")
rrect(s, M, 1.55, 5.0, 1.95, PLUM, radius=0.16)
stat_block(s, 0.9, 1.78, 4.35, "4,3%", "CV Minyak Goreng Curah, komoditas paling "
           "volatile dari 13 komoditas", big_size=48, label_color=SURFACE)
tf = tb(s, M, 3.75, 5.0, 2.9)
bullets(tf, [
    f"{lk['paling_volatil']['komoditas']} paling volatile: CV {lk['paling_volatil']['cv_pct']}%, "
    f"melonjak ke {rp(km['com_17']['maks'])}/kg pada {km['com_17']['tgl_maks']}.",
    f"{lk['paling_stabil']['komoditas']} paling stabil: CV {lk['paling_stabil']['cv_pct']}%.",
    "Pola: minyak goreng curah (tanpa merek, rantai pasok pendek) jauh lebih liar "
    "dibanding versi kemasan bermerek: kualitas kemasan = stabilitas harga.",
    "Volatilitas tinggi = ketidakpastian bagi pedagang dan konsumen.",
], size=14.5, gap=10)
bezel(s, 5.9, 1.55, 6.88, 4.6)
s.shapes.add_picture(str(CHARTS / "03_volatilitas.png"), Inches(6.22), Inches(1.93),
                     width=Inches(6.25))

# ================= Slide 7: Puncak Serentak =================
s = slide_new()
title(s, "Puncak Serentak 13 Oktober 2026")
tf = tb(s, M, 1.55, 7.35, 2.5)
bullets(tf, [
    "Tiga kelompok mencapai titik tertinggi pada 13-14 Okt 2026: "
    f"Beras {rp(kl['Beras']['indeks_maks'])}, Daging Sapi {rp(kl['Daging Sapi']['indeks_maks'])}, "
    f"Minyak Goreng {rp(MG['indeks_maks'])}. Gula Pasir menyusul tipis melampaui "
    f"puncaknya pada 21 Okt 2026: {rp(kl['Gula Pasir']['indeks_maks'])}.",
    "Pola yang sama terlihat pada proyek Harga Cabai (puncak 12 Okt 2026), "
    "ini lonjakan harga pangan luas, bukan ulah satu komoditas.",
    "Bagi pemantau kebijakan: minggu kedua Oktober 2026 adalah momen tekanan "
    "maksimum harga pangan tahun ini.",
], size=15, gap=11)
rrect(s, 8.25, 1.55, 4.53, 2.42, TANG, radius=0.16)
stat_block(s, 8.6, 1.85, 3.9, "13 Okt 2026", "tiga kelompok mencapai titik "
           "tertingginya pada 13-14 Oktober 2026; gula pasir menyusul tipis pada 21 Oktober 2026", big_color=WHITE, big_size=36,
           label_color=WHITE)
bezel(s, M, 4.2, 9.3, 2.85)
s.shapes.add_picture(str(CHARTS / "04_puncak.png"), Inches(2.55), Inches(4.28),
                     width=Inches(5.3))
rrect(s, 10.2, 4.2, 2.58, 2.85, SURFACE, radius=0.16)
tf = tb(s, 10.5, 4.58, 2.0, 1.9)
par(tf, "12 Okt 2026", 20, PLUM, bold=True, first=True, after=4)
par(tf, "puncak Harga Cabai, pola lonjakan yang sama persis.", 12.5, INK, after=0, leading=1.2)

# ================= Slide 8: Disparitas Antar Provinsi =================
s = slide_new()
title(s, "Disparitas Antar Provinsi")
b, ds, mg, gp = pv["Beras"], pv["Daging Sapi"], pv["Minyak Goreng"], pv["Gula Pasir"]
tf = tb(s, M, 1.6, 6.0, 4.15)
bullets(tf, [
    f"Beras: {b['termurah']['provinsi']} {rp(b['termurah']['rata'])}/kg vs "
    f"{b['termahal']['provinsi']} {rp(b['termahal']['rata'])}/kg, selisih {b['selisih_pct']}%.",
    f"Daging Sapi: {ds['termurah']['provinsi']} {rp(ds['termurah']['rata'])}/kg vs "
    f"{ds['termahal']['provinsi']} {rp(ds['termahal']['rata'])}/kg, selisih {ds['selisih_pct']}%.",
    f"Minyak Goreng: {mg['termurah']['provinsi']} vs {mg['termahal']['provinsi']} "
    f"(selisih {mg['selisih_pct']}%); Gula Pasir: {gp['termurah']['provinsi']} vs "
    f"{gp['termahal']['provinsi']} (selisih {gp['selisih_pct']}%).",
    "Pola konsisten: provinsi timur dan kepulauan membayar 36-41% lebih mahal, "
    "logistik dan rantai pasok adalah bagian dari masalah harga.",
], size=15, gap=11)
rrect(s, M, 5.95, 6.0, 1.0, PLUM, radius=0.14)
tf = tb(s, 0.85, 6.17, 5.5, 0.6)
par(tf, "Selisih termurah vs termahal: 36-41% di semua kelompok",
    14.5, SURFACE, bold=True, first=True, after=0)
bezel(s, 6.9, 1.55, 5.88, 5.4)
s.shapes.add_picture(str(CHARTS / "05_disparitas_provinsi.png"), Inches(7.22), Inches(1.87),
                     width=Inches(5.25))

# ================= Slide 9: Ramadan =================
s = slide_new()
title(s, "Apakah Ramadan Mendongkrak Harga?")
bezel(s, M, 1.55, 5.88, 5.4)
s.shapes.add_picture(str(CHARTS / "06_ramadan.png"), Inches(0.87), Inches(1.87),
                     width=Inches(5.25))
rb = F["ramadan"]["Beras"]
tf = tb(s, 6.85, 1.6, 5.95, 4.15)
bullets(tf, [
    f"Ramadan 2025: beras {rp(rb[0]['sebelum'])}/kg (30 hari sebelum), "
    f"{rp(rb[0]['selama'])}/kg (selama), {rp(rb[0]['sesudah'])}/kg (30 hari sesudah).",
    f"Ramadan 2026: {rp(rb[1]['sebelum'])}/kg, {rp(rb[1]['selama'])}/kg, "
    f"{rp(rb[1]['sesudah'])}/kg.",
    "Kenaikan selama Ramadan hanya 0,2-0,4%, kecil. Narasi 'harga selalu melonjak "
    "saat puasa' tidak terbukti pada data 2025-2026 untuk beras.",
    "Tekanan harga yang sesungguhnya terjadi di luar musim Ramadan: tren naik "
    "jangka panjang + lonjakan Oktober 2026.",
], size=15, gap=11)
rrect(s, 6.85, 5.95, 5.95, 1.0, TANG, radius=0.14)
tf = tb(s, 7.15, 6.17, 5.4, 0.6)
par(tf, "Efek Ramadan pada beras: hanya +0,2-0,4%", 14.5, WHITE, bold=True,
    first=True, after=0)

# ================= Slide 10: Penutup (Designed Close) =================
s = slide_new()
title(s, "Penutup")
bezel(s, M, 1.5, 7.35, 3.55)
tf = tb(s, 0.95, 1.85, 6.6, 2.95)
par(tf, "Tiga pesan utama", 16, PLUM, bold=True, first=True, after=8)
par(tf, "(1) Minyak goreng naik tertinggi (+6,8%) dan paling liar, pantau curah; "
        "(2) medio Okt 2026 adalah puncak tekanan harga pangan nasional; "
        "(3) disparitas 36-41% antarprovinsi menuntut perhatian logistik, bukan cuma produksi.",
    15, INK, after=0, leading=1.25)
rrect(s, 8.25, 1.5, 4.53, 1.68, PLUM, radius=0.16)
tf = tb(s, 8.6, 1.78, 3.85, 1.2)
par(tf, "Dashboard Power BI", 15, TANG, bold=True, first=True, after=4)
par(tf, "Star schema 5 tabel, measure DAX siap tempel, layout 4 halaman, "
        "panduan rakit langkah demi langkah.", 12.5, SURFACE, after=0, leading=1.2)
rrect(s, 8.25, 3.37, 4.53, 1.68, SURFACE, radius=0.16)
tf = tb(s, 8.6, 3.65, 3.85, 1.2)
par(tf, "Terbuka untuk direproduksi", 15, PLUM, bold=True, first=True, after=4)
par(tf, "Jalankan scripts/fetch_pihps.py lalu scripts/analyze.py: findings.json "
        "dan 12 chart dihitung ulang dari nol.", 12.5, INK, after=0, leading=1.2)
rrect(s, M, 5.35, 12.23, 1.6, PLUM, radius=0.18)
tf = tb(s, 1.0, 5.68, 7.6, 1.0)
par(tf, "BAGHIZ ASSHIDDIQIE ARASYA", 23, WHITE, bold=True, first=True, after=3)
par(tf, REPO_URL, 13.5, TANG, after=0)
tf = tb(s, 8.9, 5.82, 3.45, 0.85)
par(tf, "Terima kasih. Silakan jelajahi repo dan dashboardnya, semua angka "
        "terbuka untuk dicek ulang.", 12.5, SURFACE, first=True, after=0, leading=1.2)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print(f"PPT saved: {OUT}")
